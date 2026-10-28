#!/usr/bin/env python3
"""Build tailored resumes as .docx and export them to PDF.

  resume.py dump BASE.docx       list each paragraph's index, style and runs, to pick formatting examples
  resume.py build SPEC.json      build a resume from a spec (see references/resume.md for the format)
  resume.py pdf FILE.docx [--png]  export a PDF beside the .docx, print its page count, optionally a PNG preview

Clone mode copies formatting from paragraphs of the user's own resume. Fresh mode ("mode": "fresh")
builds a clean single-column layout. Paths in a spec are relative to the folder you run it from.
build and dump need python-docx (pip install python-docx); pdf does not.
"""
import copy
import json
import platform
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run
except ImportError:
    docx = None

# Run children that carry more than text; runs holding them are never merged or rewritten
SPECIAL = ("drawing", "pict", "object", "fldChar", "instrText", "footnoteReference", "endnoteReference")
# Markup that can sit between two runs without separating them visually
SKIPPABLE = ("bookmarkStart", "bookmarkEnd", "proofErr", "permStart", "permEnd")


def need_docx():
    if docx is None:
        sys.exit("python-docx is missing. Install it with: pip install python-docx")


def local(el):
    return el.tag.split("}")[-1]


def fmt(run):
    rpr = run._r.rPr
    return rpr.xml if rpr is not None else ""


def is_special(run):
    return any(local(c) in SPECIAL for c in run._r)


def merge_runs(p):
    """Join neighboring runs that share formatting, so runs line up with how the paragraph looks."""
    runs = list(p.runs)
    i = 0
    while i < len(runs) - 1:
        a, b = runs[i], runs[i + 1]
        between, el = [], a._r.getnext()
        while el is not None and el is not b._r and local(el) in SKIPPABLE:
            between.append(el)
            el = el.getnext()
        if el is b._r and fmt(a) == fmt(b) and not is_special(a) and not is_special(b):
            a.text = a.text + b.text
            for e in between:
                if local(e) == "proofErr":
                    e.getparent().remove(e)
            b._r.getparent().remove(b._r)
            runs.pop(i + 1)
        else:
            i += 1


def describe_run(r):
    flags = []
    if r.bold:
        flags.append("bold")
    if r.italic:
        flags.append("italic")
    if r.underline:
        flags.append("underline")
    if r.font.size:
        flags.append(f"{r.font.size.pt:g}pt")
    if r.font.name:
        flags.append(r.font.name)
    return f"[{', '.join(flags)}] " if flags else ""


def dump(path):
    need_docx()
    d = docx.Document(path)
    body = d.element.body
    tables = body.findall(qn("w:tbl"))
    if tables:
        print(f"NOTE: {len(tables)} table(s) in the body. Tables above the first rebuilt paragraph are kept; "
              "tables below it are removed on build. If the resume's content lives in tables, use fresh mode.")
    if body.xpath(".//w:txbxContent"):
        print("NOTE: text boxes found. Clone mode can't rebuild text inside them; consider fresh mode.")
    for i, p in enumerate(d.paragraphs):
        merge_runs(p)
        style = p.style.name if p.style is not None else "?"
        bullet = " bullet" if p._p.pPr is not None and p._p.pPr.numPr is not None else ""
        align = f" align={p.alignment.name}" if p.alignment is not None else ""
        print(f"[{i}] style={style}{bullet}{align}")
        if not p.runs and not p.text:
            print("    (empty)")
        for j, r in enumerate(p.runs):
            tag = " (special, kept as is)" if is_special(r) else ""
            print(f"    run {j} {describe_run(r)}{r.text!r}{tag}")
        for h in p._p.findall(qn("w:hyperlink")):
            text = "".join(t.text or "" for t in h.iter(qn("w:t")))
            print(f"    link {text!r}")


def fill(p, texts):
    """Put each text into the paragraph's runs in order; extra texts reuse the last run's formatting."""
    for h in p._p.findall(qn("w:hyperlink")):
        p._p.remove(h)
    runs = [r for r in p.runs if not is_special(r)]
    if not runs:
        for t in texts:
            p.add_run(t)
        return
    for r, t in zip(runs, texts):
        r.text = t
    for r in runs[len(texts):]:
        r._r.getparent().remove(r._r)
    last = runs[min(len(texts), len(runs)) - 1]._r
    for t in texts[len(runs):]:
        el = copy.deepcopy(last)
        last.addnext(el)
        last = el
        Run(el, p).text = t


def build_clone(spec, out):
    need_docx()
    base = Path(spec["base"])
    if not base.exists():
        sys.exit(f"base resume not found: {base}")
    d = docx.Document(base)
    body = d.element.body
    paras = d.paragraphs
    keep = int(spec.get("keep", 2))
    like = spec.get("like", {})

    def index(kind):
        if isinstance(kind, int):
            return kind
        if kind not in like:
            sys.exit(f"block kind {kind!r} has no example paragraph in 'like'")
        return int(like[kind])

    templates = {}
    for block in spec["blocks"]:
        i = index(block[0])
        if i >= len(paras):
            sys.exit(f"paragraph {i} doesn't exist; the base has {len(paras)} (run dump to check)")
        if i not in templates:
            templates[i] = copy.deepcopy(paras[i]._p)

    removed_tables = 0
    if keep < len(paras):
        start = list(body).index(paras[keep]._p)
        for el in list(body)[start:]:
            if local(el) == "p":
                body.remove(el)
            elif local(el) == "tbl":
                body.remove(el)
                removed_tables += 1
    if removed_tables:
        print(f"note: removed {removed_tables} table(s) below the kept header")

    anchor = paras[keep - 1]._p if keep > 0 else None
    for block in spec["blocks"]:
        kind, texts = block[0], [str(t) for t in block[1:]]
        el = copy.deepcopy(templates[index(kind)])
        if anchor is None:
            body.insert(0, el)
        else:
            anchor.addnext(el)
        anchor = el
        p = Paragraph(el, d._body)
        if texts:
            merge_runs(p)
            fill(p, texts)
    d.save(out)


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    line = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        line.set(qn(k), v)
    bdr.append(line)
    later = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
             "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
             "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
             "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
             "w:rPr", "w:sectPr", "w:pPrChange")
    pPr.insert_element_before(bdr, *later)


def build_fresh(spec, out):
    need_docx()
    d = docx.Document()
    for p in list(d.paragraphs):
        p._p.getparent().remove(p._p)
    margin = Inches(float(spec.get("margin", 0.6)))
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = margin
    width = sec.page_width - 2 * margin
    font, size = spec.get("font", "Calibri"), float(spec.get("size", 10.5))

    normal = d.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(size)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    pf = normal.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(0), 1.0

    def para(style=None):
        p = d.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(0)
        return p

    for block in spec["blocks"]:
        kind, texts = block[0], [str(t) for t in block[1:]]
        first = texts[0] if texts else ""
        if kind == "name":
            p = para()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(first)
            r.bold, r.font.size = True, Pt(size + 7)
        elif kind == "contact":
            p = para()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            p.add_run(" | ".join(texts)).font.size = Pt(size - 0.5)
        elif kind == "section":
            p = para()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(first.upper())
            r.bold, r.font.size = True, Pt(size + 1)
            bottom_border(p)
        elif kind in ("entry", "title"):
            p = para()
            if kind == "entry":
                p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
            r = p.add_run(first)
            r.bold, r.italic = kind == "entry", kind == "title"
            if len(texts) > 1:
                right = p.add_run("\t" + texts[1])
                right.italic = kind == "title"
        elif kind == "bullet":
            p = para("List Bullet")
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.first_line_indent = Inches(-0.15)
            p.add_run(first)
        elif kind == "skills":
            p = para()
            p.add_run(first).bold = True
            if len(texts) > 1:
                p.add_run(texts[1])
        elif kind == "text":
            para().add_run(first)
        elif kind == "blank":
            para().paragraph_format.space_after = Pt(2)
        else:
            sys.exit(f"unknown fresh block kind {kind!r}; use name, contact, section, entry, title, bullet, skills, text, blank")
    d.save(out)


def build(spec_path):
    spec = json.loads(Path(spec_path).read_text())
    out = Path(spec["out"])
    out.parent.mkdir(parents=True, exist_ok=True)
    if spec.get("base") and Path(spec["base"]).resolve() == out.resolve():
        sys.exit("'out' is the base resume; tailored versions must be new files")
    if spec.get("mode") == "fresh":
        build_fresh(spec, out)
    else:
        build_clone(spec, out)
    print(out)


def page_count(pdf):
    try:
        import logging
        from pypdf import PdfReader
        logging.getLogger("pypdf").setLevel(logging.ERROR)
        return len(PdfReader(str(pdf)).pages)
    except Exception:
        pass
    if shutil.which("pdfinfo"):
        info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
        m = re.search(r"Pages:\s+(\d+)", info)
        if m:
            return int(m.group(1))
    return len(re.findall(rb"/Type\s*/Page[^s]", Path(pdf).read_bytes()))


def osascript(script):
    return subprocess.run(["osascript", "-e", script], capture_output=True, text=True, timeout=180)


def to_pdf(path, png=False):
    path = Path(path).resolve()
    if not path.exists():
        sys.exit(f"not found: {path}")
    out = path.with_suffix(".pdf")
    # Convert in a temp folder: sandboxed Mac apps often can't write inside Desktop or Documents
    tmp = Path(tempfile.mkdtemp(prefix="scout-pdf-"))
    src = tmp / path.name
    shutil.copy(path, src)
    made = src.with_suffix(".pdf")
    tried = []

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    mac_lo = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if not soffice and mac_lo.exists():
        soffice = str(mac_lo)
    if soffice:
        tried.append("LibreOffice")
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(src)],
                       capture_output=True, timeout=180)
    if not made.exists() and platform.system() == "Darwin" and Path("/Applications/Pages.app").exists():
        tried.append("Pages")
        osascript(f'''tell application "Pages"
            set d to open POSIX file "{src}"
            export d to POSIX file "{made}" as PDF
            close d saving no
        end tell''')
    if not made.exists() and platform.system() == "Darwin" and Path("/Applications/Microsoft Word.app").exists():
        tried.append("Word")
        osascript(f'''tell application "Microsoft Word"
            open POSIX file "{src}"
            set d to active document
            save as d file name (POSIX file "{made}" as string) file format format PDF
            close d saving no
        end tell''')
    if not made.exists():
        how = ", ".join(tried) if tried else "no converter found (LibreOffice, Pages or Word)"
        sys.exit(f"couldn't make a PDF ({how}). Open {path.name} in Word or Google Docs and save it as PDF.")

    shutil.copy(made, out)
    shutil.rmtree(tmp, ignore_errors=True)
    n = page_count(out)
    print(f"{out} ({n} page{'s' if n != 1 else ''})")
    if n > 1:
        print("Over one page: cut the weakest bullet before shrinking fonts or margins.")
    if png:
        preview = out.with_suffix(".png")
        if shutil.which("pdftoppm"):
            subprocess.run(["pdftoppm", "-png", "-r", "110", "-singlefile", str(out), str(preview.with_suffix(""))],
                           capture_output=True)
        elif platform.system() == "Darwin":
            subprocess.run(["sips", "-s", "format", "png", str(out), "--out", str(preview)], capture_output=True)
        print(preview if preview.exists() else "no PNG tool found (pdftoppm or sips)")


def main():
    if len(sys.argv) < 3 or sys.argv[1] not in ("dump", "build", "pdf"):
        sys.exit(__doc__)
    cmd, target = sys.argv[1], sys.argv[2]
    if cmd == "dump":
        dump(target)
    elif cmd == "build":
        build(target)
    else:
        to_pdf(target, png="--png" in sys.argv[3:])


if __name__ == "__main__":
    main()
